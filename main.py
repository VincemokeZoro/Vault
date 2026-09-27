import tkinter as tk
from tkinter import messagebox, ttk, filedialog
from datetime import datetime, date, timedelta
import calendar
import csv
import shutil
import sqlite3
import sys
from pathlib import Path
from collections import defaultdict

# ============================================================
# VAULT
# Victorian-inspired personal finance ledger
# ============================================================

window = tk.Tk()
window.title("Vault")
window.geometry("1180x760")
window.minsize(980, 650)

# ---------- Theme ----------

LIGHT = {
    "bg": "#eee8dc",
    "sidebar": "#25221f",
    "sidebar_hover": "#3a332c",
    "sidebar_text": "#f3eadb",
    "sidebar_muted": "#b9aa94",
    "card": "#f8f3e8",
    "card_alt": "#e7dece",
    "text": "#2b2621",
    "muted": "#776d61",
    "accent": "#a67c3d",
    "accent_dark": "#765527",
    "danger": "#9a4438",
    "success": "#5e7658",
    "border": "#c9b99e",
    "today": "#d8bd83",
    "shadow": "#d1c5b3",
    "input": "#fffaf0",
}

DARK = {
    "bg": "#171513",
    "sidebar": "#0f0e0d",
    "sidebar_hover": "#2c2722",
    "sidebar_text": "#eee3cf",
    "sidebar_muted": "#8f8272",
    "card": "#211e1a",
    "card_alt": "#2c2722",
    "text": "#eee3cf",
    "muted": "#a99d8b",
    "accent": "#c39a55",
    "accent_dark": "#9b753d",
    "danger": "#c86b5e",
    "success": "#88a47c",
    "border": "#514638",
    "today": "#6e5734",
    "shadow": "#0d0c0b",
    "input": "#2a251f",
}

theme = LIGHT
dark_mode = False

# ---------- Data / Database ----------

def app_directory():
    if getattr(sys, "frozen", False):
        base = Path(sys.executable).resolve().parent
    else:
        base = Path(__file__).resolve().parent
    try:
        base.mkdir(parents=True, exist_ok=True)
        test_file = base / ".vault_write_test"
        test_file.touch(exist_ok=True)
        test_file.unlink(missing_ok=True)
        return base
    except OSError:
        fallback = Path.home() / "AppData" / "Local" / "Vault"
        fallback.mkdir(parents=True, exist_ok=True)
        return fallback


APP_DIR = app_directory()
DB_PATH = APP_DIR / "vault.db"
BACKUP_DIR = APP_DIR / "backups"

BUILT_IN_CATEGORIES = [
    "Food",
    "Transportation",
    "School",
    "Bills",
    "Gaming",
    "Personal",
    "Health",
    "Entertainment",
    "Other",
]


def get_db():
    connection = sqlite3.connect(DB_PATH)
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA foreign_keys = ON")
    return connection


def table_columns(db, table_name):
    return {row["name"] for row in db.execute(f"PRAGMA table_info({table_name})").fetchall()}


def init_db():
    with get_db() as db:
        db.execute("""
            CREATE TABLE IF NOT EXISTS transactions (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                type TEXT NOT NULL CHECK(type IN ('income', 'expense')),
                amount REAL NOT NULL,
                note TEXT NOT NULL,
                date TEXT NOT NULL,
                time TEXT NOT NULL,
                goal TEXT DEFAULT '',
                category TEXT DEFAULT 'Other',
                created_at TEXT DEFAULT ''
            )
        """)

        db.execute("""
            CREATE TABLE IF NOT EXISTS goals (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL UNIQUE,
                target REAL NOT NULL,
                saved REAL NOT NULL DEFAULT 0,
                deadline TEXT DEFAULT '',
                description TEXT DEFAULT '',
                completed INTEGER NOT NULL DEFAULT 0
            )
        """)

        db.execute("""
            CREATE TABLE IF NOT EXISTS settings (
                key TEXT PRIMARY KEY,
                value TEXT NOT NULL
            )
        """)

        db.execute("""
            CREATE TABLE IF NOT EXISTS categories (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL UNIQUE
            )
        """)

        db.execute("""
            CREATE TABLE IF NOT EXISTS recurring (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL,
                type TEXT NOT NULL CHECK(type IN ('income', 'expense')),
                amount REAL NOT NULL,
                note TEXT NOT NULL,
                category TEXT DEFAULT 'Other',
                goal TEXT DEFAULT '',
                frequency TEXT NOT NULL,
                next_date TEXT NOT NULL,
                active INTEGER NOT NULL DEFAULT 1
            )
        """)

        # Migrate databases created by the previous Vault version.
        tx_columns = table_columns(db, "transactions")
        if "category" not in tx_columns:
            db.execute("ALTER TABLE transactions ADD COLUMN category TEXT DEFAULT 'Other'")
        if "created_at" not in tx_columns:
            db.execute("ALTER TABLE transactions ADD COLUMN created_at TEXT DEFAULT ''")

        goal_columns = table_columns(db, "goals")
        if "deadline" not in goal_columns:
            db.execute("ALTER TABLE goals ADD COLUMN deadline TEXT DEFAULT ''")
        if "description" not in goal_columns:
            db.execute("ALTER TABLE goals ADD COLUMN description TEXT DEFAULT ''")
        if "completed" not in goal_columns:
            db.execute("ALTER TABLE goals ADD COLUMN completed INTEGER NOT NULL DEFAULT 0")

        for category in BUILT_IN_CATEGORIES:
            db.execute("INSERT OR IGNORE INTO categories (name) VALUES (?)", (category,))

        db.execute("""
            UPDATE transactions
            SET category = 'Other'
            WHERE category IS NULL OR TRIM(category) = ''
        """)

        db.execute("""
            UPDATE transactions
            SET created_at = date || ' ' || time
            WHERE created_at IS NULL OR TRIM(created_at) = ''
        """)


def integrity_check():
    with get_db() as db:
        result = db.execute("PRAGMA integrity_check").fetchone()[0]
    return result


def backup_database():
    try:
        if not DB_PATH.exists():
            return None
        BACKUP_DIR.mkdir(parents=True, exist_ok=True)
        today_stamp = datetime.now().strftime("%Y-%m-%d")
        existing = list(BACKUP_DIR.glob(f"vault_{today_stamp}*.db"))
        if existing:
            return existing[0]
        target = BACKUP_DIR / f"vault_{datetime.now().strftime('%Y-%m-%d_%H%M%S')}.db"
        with get_db() as db:
            db.execute("PRAGMA wal_checkpoint(PASSIVE)")
        shutil.copy2(DB_PATH, target)
        return target
    except Exception:
        return None


def load_data():
    global balance, total_saved, total_spent, transactions, goals, categories
    global recurring_items, dark_mode, theme

    transactions = []
    goals = []
    categories = []
    recurring_items = []

    with get_db() as db:
        rows = db.execute("""
            SELECT id, type, amount, note, date, time, goal, category, created_at
            FROM transactions
            ORDER BY id
        """).fetchall()

        for row in rows:
            transactions.append({
                "id": row["id"],
                "type": row["type"],
                "amount": float(row["amount"]),
                "note": row["note"],
                "date": row["date"],
                "time": row["time"],
                "goal": row["goal"] or "",
                "category": row["category"] or "Other",
                "created_at": row["created_at"] or "",
            })

        goal_rows = db.execute("""
            SELECT id, name, target, saved, deadline, description, completed
            FROM goals
            ORDER BY id
        """).fetchall()

        for row in goal_rows:
            goals.append({
                "id": row["id"],
                "name": row["name"],
                "target": float(row["target"]),
                "saved": float(row["saved"]),
                "deadline": row["deadline"] or "",
                "description": row["description"] or "",
                "completed": bool(row["completed"]),
            })

        category_rows = db.execute(
            "SELECT name FROM categories ORDER BY name"
        ).fetchall()
        categories = [row["name"] for row in category_rows]

        recurring_rows = db.execute("""
            SELECT id, name, type, amount, note, category, goal,
                   frequency, next_date, active
            FROM recurring
            ORDER BY id
        """).fetchall()

        for row in recurring_rows:
            recurring_items.append({
                "id": row["id"],
                "name": row["name"],
                "type": row["type"],
                "amount": float(row["amount"]),
                "note": row["note"],
                "category": row["category"] or "Other",
                "goal": row["goal"] or "",
                "frequency": row["frequency"],
                "next_date": row["next_date"],
                "active": bool(row["active"]),
            })

        setting = db.execute(
            "SELECT value FROM settings WHERE key = 'dark_mode'"
        ).fetchone()

    total_saved = sum(t["amount"] for t in transactions if t["type"] == "income")
    total_spent = sum(t["amount"] for t in transactions if t["type"] == "expense")
    balance = total_saved - total_spent

    dark_mode = bool(setting and setting["value"] == "1")
    theme = DARK if dark_mode else LIGHT


def reload_data_and_page():
    load_data()
    refresh_current_page()


def save_transaction(transaction):
    with get_db() as db:
        cursor = db.execute("""
            INSERT INTO transactions
                (type, amount, note, date, time, goal, category, created_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            transaction["type"],
            transaction["amount"],
            transaction["note"],
            transaction["date"],
            transaction["time"],
            transaction.get("goal", ""),
            transaction.get("category", "Other"),
            transaction.get("created_at", ""),
        ))
        return cursor.lastrowid


def update_transaction_db(transaction):
    with get_db() as db:
        db.execute("""
            UPDATE transactions
            SET type=?, amount=?, note=?, date=?, time=?,
                goal=?, category=?, created_at=?
            WHERE id=?
        """, (
            transaction["type"],
            transaction["amount"],
            transaction["note"],
            transaction["date"],
            transaction["time"],
            transaction.get("goal", ""),
            transaction.get("category", "Other"),
            transaction.get("created_at", ""),
            transaction["id"],
        ))


def delete_transaction_db(transaction_id):
    with get_db() as db:
        db.execute("DELETE FROM transactions WHERE id=?", (transaction_id,))


def save_goal(goal):
    with get_db() as db:
        cursor = db.execute("""
            INSERT INTO goals (name, target, saved, deadline, description, completed)
            VALUES (?, ?, ?, ?, ?, ?)
        """, (
            goal["name"],
            goal["target"],
            goal["saved"],
            goal.get("deadline", ""),
            goal.get("description", ""),
            1 if goal.get("completed") else 0,
        ))
        return cursor.lastrowid


def update_goal_in_db(goal):
    with get_db() as db:
        db.execute("""
            UPDATE goals
            SET name=?, target=?, saved=?, deadline=?, description=?, completed=?
            WHERE id=?
        """, (
            goal["name"],
            goal["target"],
            goal["saved"],
            goal.get("deadline", ""),
            goal.get("description", ""),
            1 if goal.get("completed") else 0,
            goal["id"],
        ))


def delete_goal_db(goal_id):
    with get_db() as db:
        db.execute("DELETE FROM goals WHERE id=?", (goal_id,))


def save_dark_mode():
    with get_db() as db:
        db.execute("""
            INSERT INTO settings (key, value) VALUES ('dark_mode', ?)
            ON CONFLICT(key) DO UPDATE SET value = excluded.value
        """, ("1" if dark_mode else "0",))


def add_category_db(name):
    with get_db() as db:
        db.execute("INSERT INTO categories (name) VALUES (?)", (name,))


def delete_category_db(name):
    if name in BUILT_IN_CATEGORIES:
        return
    with get_db() as db:
        db.execute("DELETE FROM categories WHERE name=?", (name,))


def save_recurring(item):
    with get_db() as db:
        cursor = db.execute("""
            INSERT INTO recurring
                (name, type, amount, note, category, goal, frequency, next_date, active)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            item["name"],
            item["type"],
            item["amount"],
            item["note"],
            item.get("category", "Other"),
            item.get("goal", ""),
            item["frequency"],
            item["next_date"],
            1 if item.get("active", True) else 0,
        ))
        return cursor.lastrowid


def update_recurring_db(item):
    with get_db() as db:
        db.execute("""
            UPDATE recurring
            SET name=?, type=?, amount=?, note=?, category=?, goal=?,
                frequency=?, next_date=?, active=?
            WHERE id=?
        """, (
            item["name"],
            item["type"],
            item["amount"],
            item["note"],
            item.get("category", "Other"),
            item.get("goal", ""),
            item["frequency"],
            item["next_date"],
            1 if item.get("active", True) else 0,
            item["id"],
        ))


def delete_recurring_db(item_id):
    with get_db() as db:
        db.execute("DELETE FROM recurring WHERE id=?", (item_id,))


def insert_existing_transaction_db(transaction):
    with get_db() as db:
        db.execute("""
            INSERT OR REPLACE INTO transactions
                (id, type, amount, note, date, time, goal, category, created_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            transaction["id"],
            transaction["type"],
            transaction["amount"],
            transaction["note"],
            transaction["date"],
            transaction["time"],
            transaction.get("goal", ""),
            transaction.get("category", "Other"),
            transaction.get("created_at", ""),
        ))


# ---------- State ----------

balance = 0.00
total_saved = 0.00
total_spent = 0.00
transactions = []
goals = []
categories = []
recurring_items = []

calendar_year = datetime.now().year
calendar_month = datetime.now().month
selected_date = datetime.now().strftime("%Y-%m-%d")
current_page = "Dashboard"
transaction_search = ""
transaction_type_filter = "All"
transaction_category_filter = "All"
transaction_goal_filter = "All"
transaction_sort = "Newest"
undo_stack = []


# ---------- Fonts ----------

SERIF = "Georgia"
SERIF_BOLD = "Georgia"
SANS = "Segoe UI"


# ---------- Helpers ----------

def money(value):
    return f"₱{value:,.2f}"


def clear_content():
    for widget in content.winfo_children():
        widget.destroy()


def animate_value(label, start, end, duration=400):
    steps = 20
    difference = end - start

    def tick(step=0):
        if not label.winfo_exists():
            return
        progress = step / steps
        value = start + difference * progress
        label.configure(text=money(value))
        if step < steps:
            window.after(max(10, duration // steps), lambda: tick(step + 1))

    tick()


def hover_button(button, normal, hover):
    button.bind("<Enter>", lambda e: button.configure(bg=hover))
    button.bind("<Leave>", lambda e: button.configure(bg=normal))


def add_divider(parent, pady=8):
    tk.Frame(parent, bg=theme["border"], height=1).pack(fill="x", padx=20, pady=pady)


def make_card(parent, padx=0, pady=0, **kwargs):
    card = tk.Frame(
        parent,
        bg=theme["card"],
        highlightbackground=theme["border"],
        highlightthickness=1,
        **kwargs,
    )
    if padx or pady:
        card.pack(fill="both", expand=True, padx=padx, pady=pady)
    return card


def section_title(parent, title, subtitle=None):
    row = tk.Frame(parent, bg=theme["card"])
    row.pack(fill="x", padx=20, pady=(16, 8))
    tk.Label(
        row,
        text=title,
        bg=theme["card"],
        fg=theme["text"],
        font=(SERIF_BOLD, 14, "bold"),
    ).pack(side="left")
    if subtitle:
        tk.Label(
            row,
            text=subtitle,
            bg=theme["card"],
            fg=theme["muted"],
            font=(SANS, 8),
        ).pack(side="right")
    return row


def page_header(title, subtitle=None):
    header = tk.Frame(content, bg=theme["bg"])
    header.pack(fill="x", padx=35, pady=(24, 14))

    tk.Label(
        header,
        text=title,
        bg=theme["bg"],
        fg=theme["text"],
        font=(SERIF_BOLD, 27, "bold"),
    ).pack(anchor="w")

    if subtitle:
        tk.Label(
            header,
            text=subtitle,
            bg=theme["bg"],
            fg=theme["muted"],
            font=(SANS, 10),
        ).pack(anchor="w", pady=(4, 0))

    return header


def show_page(name, builder):
    global current_page
    current_page = name
    clear_content()
    content.configure(cursor="watch")
    content.update_idletasks()
    window.after(20, lambda: finish_page(builder))


def finish_page(builder):
    if content.winfo_exists():
        content.configure(cursor="")
        builder()


def refresh_current_page():
    pages = {
        "Dashboard": show_dashboard,
        "Transactions": show_transactions,
        "Calendar": show_calendar,
        "Goals": show_goals,
        "Analytics": show_analytics,
        "Recurring": show_recurring,
        "Settings": show_settings,
    }
    pages.get(current_page, show_dashboard)()


def scrollable_frame(parent):
    outer = tk.Frame(parent, bg=theme["bg"])
    canvas = tk.Canvas(outer, bg=theme["bg"], highlightthickness=0)
    scrollbar = ttk.Scrollbar(outer, orient="vertical", command=canvas.yview)
    inner = tk.Frame(canvas, bg=theme["bg"])

    inner.bind(
        "<Configure>",
        lambda e: canvas.configure(scrollregion=canvas.bbox("all"))
    )

    canvas_window = canvas.create_window((0, 0), window=inner, anchor="nw")
    canvas.configure(yscrollcommand=scrollbar.set)

    def resize_inner(event):
        canvas.itemconfigure(canvas_window, width=event.width)

    canvas.bind("<Configure>", resize_inner)
    canvas.pack(side="left", fill="both", expand=True)
    scrollbar.pack(side="right", fill="y")
    return outer, inner


def parse_amount(text):
    return float(text.replace(",", "").replace("₱", "").strip())


def validate_date(text):
    datetime.strptime(text, "%Y-%m-%d")


def goal_names():
    return [goal["name"] for goal in goals]


def get_goal(name):
    for goal in goals:
        if goal["name"] == name:
            return goal
    return None


def adjust_goal(goal_name, amount):
    goal = get_goal(goal_name)
    if not goal:
        return
    goal["saved"] = max(0.0, goal["saved"] + amount)
    goal["saved"] = min(goal["saved"], goal["target"])
    goal["completed"] = goal["saved"] >= goal["target"]
    update_goal_in_db(goal)


def transaction_effect(transaction):
    signed = transaction["amount"]
    return signed if transaction["type"] == "income" else -signed


def current_month_range():
    now = datetime.now()
    prefix = f"{now.year:04d}-{now.month:02d}-"
    return prefix


def monthly_totals(year, month):
    prefix = f"{year:04d}-{month:02d}-"
    income = sum(
        t["amount"] for t in transactions
        if t["type"] == "income" and t["date"].startswith(prefix)
    )
    expense = sum(
        t["amount"] for t in transactions
        if t["type"] == "expense" and t["date"].startswith(prefix)
    )
    return income, expense, income - expense


def month_category_totals(year, month):
    prefix = f"{year:04d}-{month:02d}-"
    result = defaultdict(float)
    for t in transactions:
        if t["type"] == "expense" and t["date"].startswith(prefix):
            result[t.get("category", "Other")] += t["amount"]
    return dict(result)


def monthly_spent(year, month):
    return monthly_totals(year, month)[1]


def average_monthly_savings(months=3):
    values = []
    now = date.today()
    y, m = now.year, now.month
    for _ in range(months):
        inc, exp, net = monthly_totals(y, m)
        values.append(net)
        m -= 1
        if m == 0:
            y -= 1
            m = 12
    if not values:
        return 0
    return sum(values) / len(values)


def date_plus_frequency(start_date_text, frequency):
    try:
        current = datetime.strptime(start_date_text, "%Y-%m-%d").date()
    except ValueError:
        current = date.today()

    if frequency == "Daily":
        return current + timedelta(days=1)
    if frequency == "Weekly":
        return current + timedelta(days=7)
    if frequency == "Monthly":
        year = current.year + (current.month // 12)
        month = current.month % 12 + 1
        day = min(current.day, calendar.monthrange(year, month)[1])
        return date(year, month, day)
    if frequency == "Yearly":
        year = current.year + 1
        day = min(current.day, calendar.monthrange(year, current.month)[1])
        return date(year, current.month, day)
    return current


def last_backup_text():
    backups = sorted(BACKUP_DIR.glob("vault_*.db")) if BACKUP_DIR.exists() else []
    if not backups:
        return "No backup yet"
    stamp = datetime.fromtimestamp(backups[-1].stat().st_mtime)
    return stamp.strftime("%B %d, %Y at %I:%M %p")


def record_undo(action):
    undo_stack.append(action)
    if len(undo_stack) > 20:
        undo_stack.pop(0)


def undo_last():
    if not undo_stack:
        messagebox.showinfo("Undo", "There is nothing to undo.", parent=window)
        return

    action = undo_stack.pop()

    try:
        with get_db() as db:
            if action["kind"] == "add_transaction":
                db.execute("DELETE FROM transactions WHERE id=?", (action["transaction"]["id"],))
                for goal in action.get("goals_before", []):
                    db.execute(
                        "UPDATE goals SET saved=?, completed=? WHERE id=?",
                        (goal["saved"], 1 if goal["completed"] else 0, goal["id"])
                    )

            elif action["kind"] == "delete_transaction":
                t = action["transaction"]
                db.execute("""
                    INSERT OR REPLACE INTO transactions
                        (id, type, amount, note, date, time, goal, category, created_at)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                """, (
                    t["id"], t["type"], t["amount"], t["note"], t["date"],
                    t["time"], t.get("goal", ""), t.get("category", "Other"),
                    t.get("created_at", ""),
                ))
                for goal in action.get("goals_before", []):
                    db.execute(
                        "UPDATE goals SET saved=?, completed=? WHERE id=?",
                        (goal["saved"], 1 if goal["completed"] else 0, goal["id"])
                    )

            elif action["kind"] == "edit_transaction":
                old = action["old"]
                db.execute("""
                    UPDATE transactions
                    SET type=?, amount=?, note=?, date=?, time=?,
                        goal=?, category=?, created_at=?
                    WHERE id=?
                """, (
                    old["type"], old["amount"], old["note"], old["date"], old["time"],
                    old.get("goal", ""), old.get("category", "Other"),
                    old.get("created_at", ""), old["id"],
                ))
                for goal in action.get("goals_before", []):
                    db.execute(
                        "UPDATE goals SET saved=?, completed=? WHERE id=?",
                        (goal["saved"], 1 if goal["completed"] else 0, goal["id"])
                    )
        reload_data_and_page()
    except Exception as exc:
        messagebox.showerror("Undo Failed", str(exc), parent=window)


# ---------- Live clock ----------

def update_clock():
    now = datetime.now()
    clock_label.configure(
        text=now.strftime("%A, %B %d, %Y  •  %I:%M:%S %p")
    )
    window.after(1000, update_clock)


# ---------- Dashboard ----------

def show_dashboard():
    show_page("Dashboard", build_dashboard)


def build_dashboard():
    now = datetime.now()
    income, expenses, net = monthly_totals(now.year, now.month)
    category_totals = month_category_totals(now.year, now.month)

    page_header(
        "The Ledger",
        f"Good {('morning' if now.hour < 12 else 'afternoon' if now.hour < 18 else 'evening')}. "
        "Here is the state of your vault.",
    )

    balance_card = tk.Frame(
        content,
        bg=theme["card"],
        highlightbackground=theme["border"],
        highlightthickness=1,
    )
    balance_card.pack(fill="x", padx=35, pady=4)

    inner = tk.Frame(balance_card, bg=theme["card"])
    inner.pack(fill="x", padx=25, pady=17)

    left = tk.Frame(inner, bg=theme["card"])
    left.pack(side="left")

    tk.Label(
        left,
        text="CURRENT BALANCE",
        bg=theme["card"],
        fg=theme["muted"],
        font=(SANS, 9, "bold"),
    ).pack(anchor="w")

    balance_label = tk.Label(
        left,
        text=money(0),
        bg=theme["card"],
        fg=theme["text"],
        font=(SERIF_BOLD, 31, "bold"),
    )
    balance_label.pack(anchor="w", pady=(3, 0))
    animate_value(balance_label, 0, balance)

    tk.Label(
        inner,
        text="✦  PRIVATE LEDGER  ✦",
        bg=theme["card"],
        fg=theme["accent"],
        font=(SERIF, 10, "italic"),
    ).pack(side="right", padx=10)

    stats = tk.Frame(content, bg=theme["bg"])
    stats.pack(fill="x", padx=30, pady=12)

    create_stat_card(stats, "TOTAL IN", total_saved, theme["success"])
    create_stat_card(stats, "TOTAL OUT", total_spent, theme["danger"])
    create_stat_card(stats, "THIS MONTH", expenses, theme["accent"])
    create_stat_card(stats, "MONTHLY NET", net, theme["success"] if net >= 0 else theme["danger"])

    actions = tk.Frame(content, bg=theme["bg"])
    actions.pack(fill="x", padx=35, pady=(0, 10))

    add_btn = tk.Button(
        actions, text="＋  Add Money", command=add_money,
        bg=theme["accent"], fg="#fffaf0",
        activebackground=theme["accent_dark"], activeforeground="#fffaf0",
        relief="flat", font=(SANS, 10, "bold"), padx=18, pady=9,
        cursor="hand2", bd=0,
    )
    add_btn.pack(side="left", padx=(0, 7))
    hover_button(add_btn, theme["accent"], theme["accent_dark"])

    spend_btn = tk.Button(
        actions, text="−  Spend Money", command=spend_money,
        bg=theme["danger"], fg="#fffaf0",
        activebackground="#74342d", activeforeground="#fffaf0",
        relief="flat", font=(SANS, 10, "bold"), padx=18, pady=9,
        cursor="hand2", bd=0,
    )
    spend_btn.pack(side="left", padx=(0, 7))
    hover_button(spend_btn, theme["danger"], "#74342d")

    quick_btn = tk.Button(
        actions, text="⚡ Quick Add", command=quick_add,
        bg=theme["card_alt"], fg=theme["text"],
        activebackground=theme["border"], relief="flat",
        font=(SANS, 9, "bold"), padx=14, pady=9,
        cursor="hand2", bd=0,
    )
    quick_btn.pack(side="left", padx=(0, 7))

    undo_btn = tk.Button(
        actions, text="↶ Undo", command=undo_last,
        bg=theme["card_alt"], fg=theme["text"],
        activebackground=theme["border"], relief="flat",
        font=(SANS, 9, "bold"), padx=14, pady=9,
        cursor="hand2", bd=0,
    )
    undo_btn.pack(side="left")

    lower = tk.Frame(content, bg=theme["bg"])
    lower.pack(fill="both", expand=True, padx=35, pady=(0, 15))

    recent = make_card(lower)
    recent.pack(side="left", fill="both", expand=True, padx=(0, 6))

    section_title(recent, "Recent Entries", f"{len(transactions)} entries")
    add_divider(recent, 2)

    if not transactions:
        tk.Label(
            recent,
            text="The ledger is empty.\nYour first entry awaits.",
            bg=theme["card"], fg=theme["muted"],
            font=(SERIF, 11, "italic"), justify="center",
        ).pack(expand=True)
    else:
        for transaction in transactions[-7:][::-1]:
            transaction_row(recent, transaction, compact=True)

    insight = make_card(lower)
    insight.pack(side="right", fill="both", expand=True, padx=(6, 0))

    section_title(insight, "This Month", now.strftime("%B %Y"))
    add_divider(insight, 2)

    tk.Label(
        insight,
        text=f"Income     {money(income)}",
        bg=theme["card"], fg=theme["success"],
        font=(SANS, 10, "bold"),
    ).pack(anchor="w", padx=20, pady=(10, 4))

    tk.Label(
        insight,
        text=f"Spending   {money(expenses)}",
        bg=theme["card"], fg=theme["danger"],
        font=(SANS, 10, "bold"),
    ).pack(anchor="w", padx=20, pady=4)

    tk.Label(
        insight,
        text=f"Net        {money(net)}",
        bg=theme["card"], fg=theme["text"],
        font=(SANS, 10, "bold"),
    ).pack(anchor="w", padx=20, pady=4)

    top_category = max(category_totals.items(), key=lambda x: x[1], default=("None", 0))
    tk.Label(
        insight,
        text=f"Largest expense category\n{top_category[0]}  •  {money(top_category[1])}",
        bg=theme["card"], fg=theme["muted"],
        font=(SANS, 9), justify="left",
    ).pack(anchor="w", padx=20, pady=(9, 8))

    avg = average_monthly_savings()
    tk.Label(
        insight,
        text=f"3-month average net\n{money(avg)}",
        bg=theme["card"], fg=theme["muted"],
        font=(SANS, 9),
    ).pack(anchor="w", padx=20, pady=(4, 8))

    active_recurring = [r for r in recurring_items if r["active"]]
    tk.Label(
        insight,
        text=f"Recurring entries active  •  {len(active_recurring)}",
        bg=theme["card"], fg=theme["muted"],
        font=(SANS, 8),
    ).pack(anchor="w", padx=20, pady=(4, 14))


def create_stat_card(parent, title, value, accent):
    card = tk.Frame(
        parent,
        bg=theme["card"],
        highlightbackground=theme["border"],
        highlightthickness=1,
    )
    card.pack(side="left", expand=True, fill="both", padx=4)

    tk.Frame(card, bg=accent, width=4).pack(side="left", fill="y")
    body = tk.Frame(card, bg=theme["card"])
    body.pack(fill="both", expand=True, padx=13, pady=10)

    tk.Label(
        body, text=title, bg=theme["card"], fg=theme["muted"],
        font=(SANS, 8, "bold")
    ).pack(anchor="w")

    label = tk.Label(
        body, text=money(0), bg=theme["card"], fg=theme["text"],
        font=(SERIF_BOLD, 16, "bold")
    )
    label.pack(anchor="w", pady=(2, 0))
    animate_value(label, 0, value, 300)


def transaction_row(parent, transaction, compact=False):
    income = transaction["type"] == "income"
    color = theme["success"] if income else theme["danger"]
    sign = "+" if income else "−"

    row = tk.Frame(parent, bg=theme["card"])
    row.pack(fill="x", padx=18 if compact else 20, pady=4)

    date_time = f"{transaction['date']}  {transaction.get('time', '')}".strip()

    tk.Label(
        row, text=date_time, bg=theme["card"], fg=theme["muted"],
        font=(SANS, 7 if compact else 8), width=23, anchor="w"
    ).pack(side="left")

    tk.Label(
        row, text=f"{sign} {money(transaction['amount'])}",
        bg=theme["card"], fg=color,
        font=(SANS, 9 if compact else 10, "bold"),
        width=15, anchor="w"
    ).pack(side="left")

    description = transaction.get("note", "")
    if not compact:
        goal = transaction.get("goal", "")
        category = transaction.get("category", "Other")
        details = []
        if category:
            details.append(category)
        if goal:
            details.append(f"Goal: {goal}")
        if details:
            description += "  •  " + "  •  ".join(details)

    tk.Label(
        row, text=description, bg=theme["card"], fg=theme["text"],
        font=(SANS, 8 if compact else 9), anchor="w"
    ).pack(side="left", fill="x", expand=True)


# ---------- Add / Spend ----------

def transaction_dialog(kind, existing=None, quick=False):
    income = kind == "income" if existing is None else existing["type"] == "income"
    editing = existing is not None

    dialog = tk.Toplevel(window)
    dialog.title("Edit Entry" if editing else ("Quick Add" if quick else ("Add Money" if income else "Spend Money")))
    dialog.geometry("470x650" if not quick else "420x410")
    dialog.resizable(False, False)
    dialog.configure(bg=theme["bg"])
    dialog.transient(window)
    dialog.grab_set()

    tk.Label(
        dialog,
        text=("Edit Ledger Entry" if editing else
              ("Quick Entry" if quick else
               ("Add to the Vault" if income else "Record an Expense"))),
        bg=theme["bg"], fg=theme["text"],
        font=(SERIF_BOLD, 21, "bold")
    ).pack(pady=(20, 4))

    tk.Label(
        dialog,
        text="Correct the details before sealing this entry." if editing
        else "Make a small entry without opening the full ledger form." if quick
        else "Choose where this entry belongs in your ledger.",
        bg=theme["bg"], fg=theme["muted"], font=(SANS, 9)
    ).pack(pady=(0, 12))

    form = tk.Frame(
        dialog, bg=theme["card"],
        highlightbackground=theme["border"], highlightthickness=1
    )
    form.pack(fill="x", padx=28)

    def field(label, default=""):
        tk.Label(
            form, text=label, bg=theme["card"], fg=theme["muted"],
            font=(SANS, 8, "bold")
        ).pack(anchor="w", padx=18, pady=(9, 3))

        entry = tk.Entry(
            form, bg=theme["input"], fg=theme["text"],
            insertbackground=theme["text"], relief="flat",
            highlightbackground=theme["border"], highlightthickness=1,
            font=(SANS, 10)
        )
        entry.pack(fill="x", padx=18, pady=(0, 2), ipady=5)
        if default:
            entry.insert(0, default)
        return entry

    default_amount = "" if existing is None else str(existing["amount"])
    default_note = (
        "Added money" if income else "Expense"
    ) if existing is None else existing["note"]
    default_date = datetime.now().strftime("%Y-%m-%d") if existing is None else existing["date"]
    default_time = datetime.now().strftime("%I:%M:%S %p") if existing is None else existing.get("time", "")

    if quick:
        tk.Label(
            form, text="TYPE", bg=theme["card"], fg=theme["muted"],
            font=(SANS, 8, "bold")
        ).pack(anchor="w", padx=18, pady=(10, 3))
        quick_type_var = tk.StringVar(value="Income")
        quick_type_menu = ttk.Combobox(
            form, textvariable=quick_type_var,
            values=["Income", "Expense"], state="readonly", font=(SANS, 9)
        )
        quick_type_menu.pack(fill="x", padx=18, pady=(0, 5), ipady=4)
    else:
        quick_type_var = None
        quick_type_menu = None

    amount_entry = field("AMOUNT", default_amount)
    note_entry = field("NOTE", default_note)

    if not quick:
        date_entry = field("DATE", default_date)
        time_entry = field("TIME", default_time)

        tk.Label(
            form, text="CATEGORY", bg=theme["card"], fg=theme["muted"],
            font=(SANS, 8, "bold")
        ).pack(anchor="w", padx=18, pady=(7, 3))

        category_values = categories or BUILT_IN_CATEGORIES
        category_var = tk.StringVar(
            value=(existing.get("category", "Other") if existing else "Other")
        )
        category_menu = ttk.Combobox(
            form, textvariable=category_var,
            values=category_values, state="readonly", font=(SANS, 9)
        )
        category_menu.pack(fill="x", padx=18, pady=(0, 5), ipady=4)

        tk.Label(
            form, text="GOAL", bg=theme["card"], fg=theme["muted"],
            font=(SANS, 8, "bold")
        ).pack(anchor="w", padx=18, pady=(7, 3))

        goal_values = ["None"] + goal_names()
        existing_goal = existing.get("goal", "") if existing else ""
        goal_var = tk.StringVar(value=existing_goal or "None")
        goal_menu = ttk.Combobox(
            form, textvariable=goal_var,
            values=goal_values, state="readonly", font=(SANS, 9)
        )
        goal_menu.pack(fill="x", padx=18, pady=(0, 12), ipady=4)
    else:
        date_entry = None
        time_entry = None
        category_var = tk.StringVar(value="Other")
        goal_var = tk.StringVar(value="None")

    tk.Label(
        dialog,
        text="Ctrl+Enter seals the entry   •   Esc cancels",
        bg=theme["bg"], fg=theme["muted"],
        font=(SANS, 8)
    ).pack(pady=(8, 5))

    def submit():
        try:
            selected_income = income
            if quick:
                selected_income = quick_type_var.get() == "Income"

            amount = parse_amount(amount_entry.get())
            if amount <= 0:
                raise ValueError

            note = note_entry.get().strip() or (
                "Added money" if selected_income else "Expense"
            )

            if quick:
                entry_date = datetime.now().strftime("%Y-%m-%d")
                entry_time = datetime.now().strftime("%I:%M:%S %p")
                chosen_category = "Other"
                chosen_goal = ""
            else:
                entry_date = date_entry.get().strip()
                entry_time = time_entry.get().strip()
                validate_date(entry_date)
                if not entry_time:
                    raise ValueError
                chosen_category = category_var.get().strip() or "Other"
                chosen_goal = goal_var.get()
                if chosen_goal == "None":
                    chosen_goal = ""

            if editing:
                old_balance = balance
                old_effect = transaction_effect(existing)
                prospective_balance = old_balance - old_effect
                prospective_balance += amount if selected_income else -amount
                if prospective_balance < -1e-9:
                    messagebox.showerror(
                        "Insufficient Balance",
                        "That correction would make the recorded balance negative.",
                        parent=dialog
                    )
                    return
            elif not selected_income and amount > balance:
                messagebox.showerror(
                    "Insufficient Balance",
                    "You don't have enough money recorded.",
                    parent=dialog
                )
                return

            if editing:
                old = dict(existing)
                goals_before = []
                names = set(filter(None, [old.get("goal", ""), chosen_goal]))
                for goal in goals:
                    if goal["name"] in names:
                        goals_before.append(dict(goal))

                old_effect = transaction_effect(old)
                if old.get("goal"):
                    adjust_goal(old["goal"], -old_effect)

                updated = dict(old)
                updated.update({
                    "type": "income" if selected_income else "expense",
                    "amount": amount,
                    "note": note,
                    "date": entry_date,
                    "time": entry_time,
                    "goal": chosen_goal,
                    "category": chosen_category,
                })

                new_effect = transaction_effect(updated)
                if chosen_goal:
                    adjust_goal(chosen_goal, new_effect)

                update_transaction_db(updated)
                record_undo({
                    "kind": "edit_transaction",
                    "old": old,
                    "goals_before": goals_before,
                })
            else:
                entry = {
                    "type": "income" if selected_income else "expense",
                    "amount": amount,
                    "note": note,
                    "date": entry_date,
                    "time": entry_time,
                    "goal": chosen_goal,
                    "category": chosen_category,
                    "created_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                }

                old_goals = [dict(g) for g in goals if g["name"] == chosen_goal] if chosen_goal else []
                new_id = save_transaction(entry)
                entry["id"] = new_id

                if chosen_goal:
                    adjust_goal(chosen_goal, transaction_effect(entry))

                record_undo({
                    "kind": "add_transaction",
                    "transaction": dict(entry),
                    "goals_before": old_goals,
                })

            dialog.destroy()
            reload_data_and_page()

        except ValueError:
            messagebox.showerror(
                "Invalid Entry",
                "Enter a positive amount and a valid date using YYYY-MM-DD.",
                parent=dialog
            )

    button = tk.Button(
        dialog,
        text="Save Changes" if editing else "Seal Entry",
        command=submit,
        bg=theme["accent"] if income else theme["danger"],
        fg="#fffaf0",
        activebackground=theme["accent_dark"] if income else "#74342d",
        activeforeground="#fffaf0",
        relief="flat", bd=0,
        font=(SANS, 10, "bold"), padx=28, pady=10, cursor="hand2"
    )
    button.pack(pady=(5, 12))

    def refresh_quick_button(*_):
        if not quick:
            return
        selected = quick_type_var.get() == "Income"
        bg = theme["accent"] if selected else theme["danger"]
        hover = theme["accent_dark"] if selected else "#74342d"
        button.configure(bg=bg, activebackground=hover)

    if quick:
        quick_type_var.trace_add("write", refresh_quick_button)
        refresh_quick_button()

    hover_button(
        button,
        theme["accent"] if income else theme["danger"],
        theme["accent_dark"] if income else "#74342d"
    )

    dialog.bind("<Control-Return>", lambda e: submit())
    dialog.bind("<Escape>", lambda e: dialog.destroy())
    amount_entry.focus_set()


def add_money():
    transaction_dialog("income")


def spend_money():
    transaction_dialog("expense")


def quick_add():
    transaction_dialog("income", quick=True)


# ---------- Transactions ----------

def show_transactions():
    show_page("Transactions", build_transactions)


def filtered_transactions():
    result = list(transactions)
    query = transaction_search.lower().strip()

    if query:
        result = [
            t for t in result
            if query in " ".join([
                t.get("note", ""),
                t.get("category", ""),
                t.get("goal", ""),
                t.get("date", ""),
            ]).lower()
        ]

    if transaction_type_filter != "All":
        wanted = "income" if transaction_type_filter == "Income" else "expense"
        result = [t for t in result if t["type"] == wanted]

    if transaction_category_filter != "All":
        result = [
            t for t in result
            if t.get("category", "Other") == transaction_category_filter
        ]

    if transaction_goal_filter != "All":
        result = [
            t for t in result
            if (t.get("goal") or "None") == transaction_goal_filter
        ]

    if transaction_sort == "Newest":
        result.sort(key=lambda t: (t["date"], t.get("time", ""), t["id"]), reverse=True)
    elif transaction_sort == "Oldest":
        result.sort(key=lambda t: (t["date"], t.get("time", ""), t["id"]))
    elif transaction_sort == "Highest Amount":
        result.sort(key=lambda t: t["amount"], reverse=True)
    elif transaction_sort == "Lowest Amount":
        result.sort(key=lambda t: t["amount"])

    return result


def build_transactions():
    global transaction_search, transaction_type_filter
    global transaction_category_filter, transaction_goal_filter, transaction_sort

    page_header(
        "Ledger Entries",
        "Search, filter, correct, or remove anything in the record.",
    )

    toolbar = tk.Frame(content, bg=theme["bg"])
    toolbar.pack(fill="x", padx=35, pady=(0, 8))

    search_var = tk.StringVar(value=transaction_search)
    search_entry = tk.Entry(
        toolbar, textvariable=search_var,
        bg=theme["input"], fg=theme["text"], insertbackground=theme["text"],
        relief="flat", highlightbackground=theme["border"], highlightthickness=1,
        font=(SANS, 9)
    )
    search_entry.pack(side="left", fill="x", expand=True, ipady=7, padx=(0, 6))
    search_entry.insert(0, "")
    search_entry.focus_set()

    type_var = tk.StringVar(value=transaction_type_filter)
    type_menu = ttk.Combobox(
        toolbar, textvariable=type_var,
        values=["All", "Income", "Expense"], state="readonly", width=10
    )
    type_menu.pack(side="left", padx=3)

    category_var = tk.StringVar(value=transaction_category_filter)
    category_menu = ttk.Combobox(
        toolbar, textvariable=category_var,
        values=["All"] + (categories or BUILT_IN_CATEGORIES),
        state="readonly", width=13
    )
    category_menu.pack(side="left", padx=3)

    goal_var = tk.StringVar(value=transaction_goal_filter)
    goal_menu = ttk.Combobox(
        toolbar, textvariable=goal_var,
        values=["All", "None"] + goal_names(), state="readonly", width=15
    )
    goal_menu.pack(side="left", padx=3)

    sort_var = tk.StringVar(value=transaction_sort)
    sort_menu = ttk.Combobox(
        toolbar, textvariable=sort_var,
        values=["Newest", "Oldest", "Highest Amount", "Lowest Amount"],
        state="readonly", width=15
    )
    sort_menu.pack(side="left", padx=3)

    card = make_card(content)
    card.pack(fill="both", expand=True, padx=35, pady=(0, 10))

    header = tk.Frame(card, bg=theme["card"])
    header.pack(fill="x", padx=16, pady=10)

    headers = [
        ("DATE / TIME", 21),
        ("AMOUNT", 16),
        ("CATEGORY", 14),
        ("GOAL", 17),
        ("NOTE", 25),
    ]
    for label, width in headers:
        tk.Label(
            header, text=label, width=width,
            bg=theme["card"], fg=theme["muted"],
            font=(SANS, 8, "bold"), anchor="w"
        ).pack(side="left")

    list_frame = tk.Frame(card, bg=theme["card"])
    list_frame.pack(fill="both", expand=True, padx=6, pady=(0, 8))

    scrollbar = ttk.Scrollbar(list_frame, orient="vertical")
    list_canvas = tk.Canvas(list_frame, bg=theme["card"], highlightthickness=0)
    rows_frame = tk.Frame(list_canvas, bg=theme["card"])

    rows_frame.bind(
        "<Configure>",
        lambda e: list_canvas.configure(scrollregion=list_canvas.bbox("all"))
    )
    canvas_window = list_canvas.create_window((0, 0), window=rows_frame, anchor="nw")
    list_canvas.configure(yscrollcommand=scrollbar.set)

    def resize_rows(event):
        list_canvas.itemconfigure(canvas_window, width=event.width)

    list_canvas.bind("<Configure>", resize_rows)
    list_canvas.pack(side="left", fill="both", expand=True)
    scrollbar.pack(side="right", fill="y")

    def render_rows(*_):
        global transaction_search, transaction_type_filter
        global transaction_category_filter, transaction_goal_filter, transaction_sort

        transaction_search = search_var.get()
        transaction_type_filter = type_var.get()
        transaction_category_filter = category_var.get()
        transaction_goal_filter = goal_var.get()
        transaction_sort = sort_var.get()

        for child in rows_frame.winfo_children():
            child.destroy()

        filtered = filtered_transactions()

        if not filtered:
            tk.Label(
                rows_frame, text="No entries match these filters.",
                bg=theme["card"], fg=theme["muted"],
                font=(SERIF, 11, "italic")
            ).pack(pady=45)
            return

        for t in filtered:
            add_transaction_table_row(rows_frame, t)

    for widget in (search_entry, type_menu, category_menu, goal_menu, sort_menu):
        if isinstance(widget, ttk.Combobox):
            widget.bind("<<ComboboxSelected>>", render_rows)
        else:
            widget.bind("<KeyRelease>", render_rows)

    action_bar = tk.Frame(content, bg=theme["bg"])
    action_bar.pack(fill="x", padx=35, pady=(0, 12))

    tk.Label(
        action_bar, text="Double-click an entry to edit it.",
        bg=theme["bg"], fg=theme["muted"], font=(SANS, 8)
    ).pack(side="left")

    undo_btn = tk.Button(
        action_bar, text="↶ Undo", command=undo_last,
        bg=theme["card_alt"], fg=theme["text"],
        activebackground=theme["border"], relief="flat", bd=0,
        font=(SANS, 9, "bold"), padx=15, pady=7, cursor="hand2"
    )
    undo_btn.pack(side="right", padx=4)

    render_rows()


def add_transaction_table_row(parent, t):
    income = t["type"] == "income"
    color = theme["success"] if income else theme["danger"]
    sign = "+" if income else "−"

    row = tk.Frame(
        parent,
        bg=theme["card"],
        highlightbackground=theme["border"],
        highlightthickness=1,
    )
    row.pack(fill="x", pady=3)

    labels = [
        (f"{t['date']}  {t.get('time', '')}", 20, theme["muted"], (SANS, 8, "normal")),
        (f"{sign} {money(t['amount'])}", 15, color, (SANS, 9, "bold")),
        (t.get("category", "Other"), 13, theme["text"], (SANS, 8, "normal")),
        (t.get("goal", "") or "None", 15, theme["muted"], (SANS, 8, "normal")),
    ]

    for text_value, width, fg, font in labels:
        tk.Label(
            row, text=text_value, width=width,
            bg=theme["card"], fg=fg, font=font, anchor="w"
        ).pack(side="left", padx=(8, 0))

    tk.Label(
        row, text=t.get("note", ""), bg=theme["card"],
        fg=theme["text"], font=(SANS, 8), anchor="w"
    ).pack(side="left", fill="x", expand=True, padx=(7, 3))

    edit_btn = tk.Button(
        row, text="Edit", command=lambda x=t: edit_transaction(x),
        bg=theme["card_alt"], fg=theme["text"],
        activebackground=theme["border"], relief="flat", bd=0,
        font=(SANS, 7, "bold"), padx=8, pady=4, cursor="hand2"
    )
    edit_btn.pack(side="left", padx=2)

    delete_btn = tk.Button(
        row, text="Delete", command=lambda x=t: delete_transaction(x),
        bg=theme["card_alt"], fg=theme["danger"],
        activebackground=theme["border"], relief="flat", bd=0,
        font=(SANS, 7, "bold"), padx=8, pady=4, cursor="hand2"
    )
    delete_btn.pack(side="left", padx=(2, 6))

    def edit(_=None):
        edit_transaction(t)

    def on_enter(_):
        row.configure(bg=theme["card_alt"])
        for child in row.winfo_children():
            child.configure(bg=theme["card_alt"])

    def on_leave(_):
        row.configure(bg=theme["card"])
        for child in row.winfo_children():
            child.configure(bg=theme["card"])

    for widget in [row] + list(row.winfo_children()):
        widget.bind("<Double-Button-1>", edit)
        widget.bind("<Enter>", on_enter)
        widget.bind("<Leave>", on_leave)


def edit_transaction(transaction):
    transaction_dialog(transaction["type"], existing=transaction)


def delete_transaction(transaction):
    confirm = messagebox.askyesno(
        "Delete Entry",
        f"Delete {money(transaction['amount'])} recorded on {transaction['date']}?\n\n"
        "The balance and linked goal will be corrected automatically.",
        parent=window
    )
    if not confirm:
        return

    goals_before = [dict(g) for g in goals if g["name"] == transaction.get("goal")]
    delete_transaction_db(transaction["id"])

    if transaction.get("goal"):
        goal = get_goal(transaction["goal"])
        if goal:
            adjust_goal(transaction["goal"], -transaction_effect(transaction))

    record_undo({
        "kind": "delete_transaction",
        "transaction": dict(transaction),
        "goals_before": goals_before,
    })
    reload_data_and_page()


# ---------- Calendar ----------

def change_month(delta):
    global calendar_month, calendar_year
    calendar_month += delta
    if calendar_month < 1:
        calendar_month = 12
        calendar_year -= 1
    elif calendar_month > 12:
        calendar_month = 1
        calendar_year += 1
    show_calendar()


def jump_today():
    global calendar_year, calendar_month, selected_date
    now = datetime.now()
    calendar_year = now.year
    calendar_month = now.month
    selected_date = now.strftime("%Y-%m-%d")
    show_calendar()


def select_calendar_day(day):
    global selected_date
    if day == 0:
        return
    selected_date = f"{calendar_year:04d}-{calendar_month:02d}-{day:02d}"
    show_calendar()


def day_total(date_text):
    total = 0
    for t in transactions:
        if t["date"] == date_text:
            total += transaction_effect(t)
    return total


def show_calendar():
    show_page("Calendar", build_calendar)


def build_calendar():
    now = datetime.now()
    page_header(
        "The Calendar",
        "Follow the movement of your money through time.",
    )

    top = tk.Frame(content, bg=theme["bg"])
    top.pack(fill="x", padx=35, pady=(0, 10))

    left = tk.Frame(top, bg=theme["bg"])
    left.pack(side="left")

    prev_btn = tk.Button(
        left, text="‹", command=lambda: change_month(-1),
        bg=theme["card_alt"], fg=theme["text"],
        activebackground=theme["border"], relief="flat", bd=0,
        font=(SERIF_BOLD, 18, "bold"), width=3, cursor="hand2"
    )
    prev_btn.pack(side="left", padx=(0, 6))

    tk.Label(
        left, text=f"{calendar.month_name[calendar_month].upper()} {calendar_year}",
        bg=theme["bg"], fg=theme["text"],
        font=(SERIF_BOLD, 18, "bold")
    ).pack(side="left", padx=8)

    next_btn = tk.Button(
        left, text="›", command=lambda: change_month(1),
        bg=theme["card_alt"], fg=theme["text"],
        activebackground=theme["border"], relief="flat", bd=0,
        font=(SERIF_BOLD, 18, "bold"), width=3, cursor="hand2"
    )
    next_btn.pack(side="left", padx=6)

    today_btn = tk.Button(
        top, text="TODAY", command=jump_today,
        bg=theme["accent"], fg="#fffaf0",
        activebackground=theme["accent_dark"], relief="flat", bd=0,
        font=(SANS, 8, "bold"), padx=13, pady=7, cursor="hand2"
    )
    today_btn.pack(side="right")

    calendar_area = tk.Frame(
        content, bg=theme["card"],
        highlightbackground=theme["border"], highlightthickness=1
    )
    calendar_area.pack(fill="both", expand=True, padx=35, pady=(0, 8))

    weekdays = ["MON", "TUE", "WED", "THU", "FRI", "SAT", "SUN"]
    for col, day_name in enumerate(weekdays):
        tk.Label(
            calendar_area, text=day_name,
            bg=theme["card"], fg=theme["muted"], font=(SANS, 8, "bold")
        ).grid(row=0, column=col, sticky="ew", pady=7)
        calendar_area.grid_columnconfigure(col, weight=1)

    month_grid = calendar.monthcalendar(calendar_year, calendar_month)
    for row_index, week in enumerate(month_grid, start=1):
        calendar_area.grid_rowconfigure(row_index, weight=1)
        for col, day in enumerate(week):
            if day == 0:
                tk.Frame(calendar_area, bg=theme["card"]).grid(
                    row=row_index, column=col, sticky="nsew", padx=2, pady=2
                )
                continue

            date_text = f"{calendar_year:04d}-{calendar_month:02d}-{day:02d}"
            is_today = date_text == now.strftime("%Y-%m-%d")
            is_selected = date_text == selected_date
            total = day_total(date_text)

            bg = theme["today"] if is_selected else (
                theme["card_alt"] if is_today else theme["card"]
            )

            cell = tk.Frame(
                calendar_area, bg=bg,
                highlightbackground=theme["border"], highlightthickness=1,
                cursor="hand2"
            )
            cell.grid(
                row=row_index, column=col, sticky="nsew",
                padx=2, pady=2
            )

            number = tk.Label(
                cell, text=str(day), bg=bg, fg=theme["text"],
                font=(SERIF_BOLD, 13, "bold"), cursor="hand2"
            )
            number.pack(anchor="nw", padx=7, pady=(4, 0))

            if is_today:
                tk.Label(
                    cell, text="TODAY", bg=bg, fg=theme["accent_dark"],
                    font=(SANS, 6, "bold"), cursor="hand2"
                ).pack(anchor="nw", padx=7)

            if total != 0:
                color = theme["success"] if total > 0 else theme["danger"]
                sign = "+" if total > 0 else "−"
                tk.Label(
                    cell, text=f"{sign}{money(abs(total))}",
                    bg=bg, fg=color, font=(SANS, 7, "bold"),
                    cursor="hand2"
                ).pack(anchor="sw", padx=7, pady=(2, 4))

            for widget in (cell, number):
                widget.bind("<Button-1>", lambda e, d=day: select_calendar_day(d))

    details = tk.Frame(
        content, bg=theme["card"],
        highlightbackground=theme["border"], highlightthickness=1
    )
    details.pack(fill="x", padx=35, pady=(0, 12))

    try:
        selected_dt = datetime.strptime(selected_date, "%Y-%m-%d")
        selected_title = selected_dt.strftime("%A, %B %d, %Y")
    except ValueError:
        selected_title = selected_date

    tk.Label(
        details, text=selected_title, bg=theme["card"], fg=theme["text"],
        font=(SERIF_BOLD, 13, "bold")
    ).pack(anchor="w", padx=18, pady=(9, 2))

    selected_transactions = [
        t for t in transactions if t["date"] == selected_date
    ]

    if not selected_transactions:
        tk.Label(
            details, text="No transactions recorded for this day.",
            bg=theme["card"], fg=theme["muted"], font=(SANS, 9)
        ).pack(anchor="w", padx=18, pady=(0, 9))
    else:
        for t in selected_transactions:
            color = theme["success"] if t["type"] == "income" else theme["danger"]
            sign = "+" if t["type"] == "income" else "−"
            goal_text = f"   •   Goal: {t.get('goal')}" if t.get("goal") else ""
            category_text = f"   •   {t.get('category', 'Other')}"
            tk.Label(
                details,
                text=f"{t.get('time','')}   {sign}{money(t['amount'])}   "
                     f"{t['note']}{category_text}{goal_text}",
                bg=theme["card"], fg=color, font=(SANS, 8, "bold")
            ).pack(anchor="w", padx=18, pady=2)

        tk.Label(
            details, text=f"Day balance: {money(day_total(selected_date))}",
            bg=theme["card"], fg=theme["text"],
            font=(SANS, 9, "bold")
        ).pack(anchor="w", padx=18, pady=(3, 9))


# ---------- Goals ----------

def goal_dialog(existing=None):
    editing = existing is not None
    dialog = tk.Toplevel(window)
    dialog.title("Edit Goal" if editing else "New Savings Goal")
    dialog.geometry("450x500")
    dialog.resizable(False, False)
    dialog.configure(bg=theme["bg"])
    dialog.transient(window)
    dialog.grab_set()

    tk.Label(
        dialog,
        text="Edit Savings Goal" if editing else "New Savings Goal",
        bg=theme["bg"], fg=theme["text"],
        font=(SERIF_BOLD, 21, "bold")
    ).pack(pady=(20, 4))

    tk.Label(
        dialog,
        text="Give your next objective a proper place in the ledger.",
        bg=theme["bg"], fg=theme["muted"], font=(SANS, 9)
    ).pack(pady=(0, 12))

    form = tk.Frame(
        dialog, bg=theme["card"],
        highlightbackground=theme["border"], highlightthickness=1
    )
    form.pack(fill="x", padx=28)

    def labeled_entry(label, value=""):
        tk.Label(
            form, text=label, bg=theme["card"], fg=theme["muted"],
            font=(SANS, 8, "bold")
        ).pack(anchor="w", padx=18, pady=(10, 3))
        entry = tk.Entry(
            form, bg=theme["input"], fg=theme["text"],
            insertbackground=theme["text"], relief="flat",
            highlightbackground=theme["border"], highlightthickness=1,
            font=(SANS, 9)
        )
        entry.pack(fill="x", padx=18, ipady=5)
        if value:
            entry.insert(0, value)
        return entry

    name_entry = labeled_entry("GOAL NAME", existing["name"] if editing else "")
    target_entry = labeled_entry(
        "TARGET AMOUNT", str(existing["target"]) if editing else ""
    )
    deadline_entry = labeled_entry(
        "DEADLINE  (YYYY-MM-DD, optional)",
        existing.get("deadline", "") if editing else ""
    )
    description_entry = labeled_entry(
        "DESCRIPTION", existing.get("description", "") if editing else ""
    )

    tk.Label(
        dialog,
        text="Current allocation: " + money(existing["saved"]) if editing else
             "Money is allocated through ledger transactions.",
        bg=theme["bg"], fg=theme["muted"], font=(SANS, 8)
    ).pack(pady=9)

    def submit():
        try:
            name = name_entry.get().strip()
            target = parse_amount(target_entry.get())
            deadline = deadline_entry.get().strip()
            description = description_entry.get().strip()

            if not name or target <= 0:
                raise ValueError

            if deadline:
                validate_date(deadline)

            if editing:
                old_name = existing["name"]
                before = dict(existing)

                new_saved = min(existing["saved"], target)
                completed = new_saved >= target

                with get_db() as db:
                    db.execute("""
                        UPDATE goals
                        SET name=?, target=?, deadline=?, description=?, completed=?, saved=?
                        WHERE id=?
                    """, (
                        name, target, deadline, description,
                        1 if completed else 0, new_saved, existing["id"]
                    ))
                    if old_name != name:
                        db.execute(
                            "UPDATE transactions SET goal=? WHERE goal=?",
                            (name, old_name)
                        )

                load_data()
            else:
                new_goal = {
                    "name": name,
                    "target": target,
                    "saved": 0.0,
                    "deadline": deadline,
                    "description": description,
                    "completed": False,
                }
                save_goal(new_goal)
                load_data()

            dialog.destroy()
            show_goals()

        except sqlite3.IntegrityError:
            messagebox.showerror(
                "Duplicate Goal",
                "A goal with that name already exists.",
                parent=dialog
            )
        except ValueError:
            messagebox.showerror(
                "Invalid Goal",
                "Enter a goal name, positive target, and valid optional deadline.",
                parent=dialog
            )

    btn = tk.Button(
        dialog,
        text="Save Goal" if editing else "Create Goal",
        command=submit,
        bg=theme["accent"], fg="#fffaf0",
        activebackground=theme["accent_dark"], relief="flat", bd=0,
        font=(SANS, 10, "bold"), padx=24, pady=9, cursor="hand2"
    )
    btn.pack(pady=12)
    hover_button(btn, theme["accent"], theme["accent_dark"])
    dialog.bind("<Control-Return>", lambda e: submit())
    dialog.bind("<Escape>", lambda e: dialog.destroy())


def add_goal():
    goal_dialog()


def edit_goal(goal):
    goal_dialog(existing=goal)


def delete_goal(goal):
    confirm = messagebox.askyesno(
        "Delete Goal",
        f"Delete the goal “{goal['name']}”?\n\n"
        "Existing transactions will be kept, but their goal link will be cleared.",
        parent=window
    )
    if not confirm:
        return

    with get_db() as db:
        db.execute("UPDATE transactions SET goal='' WHERE goal=?", (goal["name"],))
        db.execute("DELETE FROM goals WHERE id=?", (goal["id"],))
    reload_data_and_page()


def show_goals():
    show_page("Goals", build_goals)


def days_until(deadline):
    try:
        return (datetime.strptime(deadline, "%Y-%m-%d").date() - date.today()).days
    except (ValueError, TypeError):
        return None


def build_goals():
    header = page_header(
        "Savings Goals",
        "Give your money somewhere meaningful to go.",
    )

    new_btn = tk.Button(
        header, text="＋ New Goal", command=add_goal,
        bg=theme["accent"], fg="#fffaf0",
        activebackground=theme["accent_dark"], relief="flat", bd=0,
        font=(SANS, 9, "bold"), padx=15, pady=8, cursor="hand2"
    )
    new_btn.pack(side="right", anchor="n")
    hover_button(new_btn, theme["accent"], theme["accent_dark"])

    outer, goals_frame = scrollable_frame(content)
    outer.pack(fill="both", expand=True, padx=35, pady=5)

    if not goals:
        empty = make_card(goals_frame)
        empty.pack(fill="x")
        tk.Label(
            empty, text="No savings goals yet.",
            bg=theme["card"], fg=theme["muted"],
            font=(SERIF, 12, "italic")
        ).pack(pady=(25, 4))
        tk.Label(
            empty, text="Create one for something you are building toward.",
            bg=theme["card"], fg=theme["muted"], font=(SANS, 9)
        ).pack(pady=(0, 25))
        return

    for goal in goals:
        card = tk.Frame(
            goals_frame, bg=theme["card"],
            highlightbackground=theme["border"], highlightthickness=1
        )
        card.pack(fill="x", pady=6)

        saved = goal["saved"]
        target = goal["target"]
        percentage = min((saved / target) * 100, 100) if target else 0
        remaining = max(target - saved, 0)

        title_row = tk.Frame(card, bg=theme["card"])
        title_row.pack(fill="x", padx=20, pady=(13, 2))

        tk.Label(
            title_row, text=goal["name"],
            bg=theme["card"], fg=theme["text"],
            font=(SERIF_BOLD, 14, "bold")
        ).pack(side="left")

        status = "FULFILLED" if goal["completed"] else "IN PROGRESS"
        status_color = theme["success"] if goal["completed"] else theme["accent"]

        tk.Label(
            title_row, text=status,
            bg=theme["card"], fg=status_color,
            font=(SANS, 7, "bold")
        ).pack(side="right", padx=(5, 0))

        tk.Label(
            card,
            text=f"{money(saved)} / {money(target)}    •    {percentage:.1f}%",
            bg=theme["card"], fg=theme["muted"], font=(SANS, 9)
        ).pack(anchor="w", padx=20)

        bar_bg = tk.Frame(card, bg=theme["card_alt"], height=9)
        bar_bg.pack(fill="x", padx=20, pady=10)
        bar_bg.pack_propagate(False)

        if percentage > 0:
            fill = tk.Frame(bar_bg, bg=theme["success"] if goal["completed"] else theme["accent"])
            fill.place(relwidth=percentage / 100, relheight=1)

        deadline = goal.get("deadline", "")
        if deadline:
            left_days = days_until(deadline)
            deadline_text = (
                f"Deadline: {deadline}  •  {left_days} days left"
                if left_days is not None and left_days >= 0
                else f"Deadline: {deadline}  •  overdue"
            )
        else:
            deadline_text = "No deadline set"

        avg = average_monthly_savings()
        months_needed = (remaining / avg) if avg > 0 and remaining > 0 else None
        pace_text = (
            f"At your 3-month average, about {months_needed:.1f} months"
            if months_needed is not None else
            "No positive savings pace calculated yet"
        )

        info = tk.Frame(card, bg=theme["card"])
        info.pack(fill="x", padx=20)

        tk.Label(
            info, text=f"Remaining: {money(remaining)}",
            bg=theme["card"], fg=theme["text"], font=(SANS, 9, "bold")
        ).pack(side="left")

        tk.Label(
            info, text=f"  •  {deadline_text}",
            bg=theme["card"], fg=theme["muted"], font=(SANS, 8)
        ).pack(side="left")

        tk.Label(
            card, text=f"Suggested pace: {pace_text}",
            bg=theme["card"], fg=theme["muted"], font=(SANS, 8)
        ).pack(anchor="w", padx=20, pady=(4, 0))

        if goal.get("description"):
            tk.Label(
                card, text=goal["description"],
                bg=theme["card"], fg=theme["muted"], font=(SANS, 8),
                wraplength=800, justify="left"
            ).pack(anchor="w", padx=20, pady=(3, 0))

        actions = tk.Frame(card, bg=theme["card"])
        actions.pack(fill="x", padx=20, pady=(8, 13))

        edit_btn = tk.Button(
            actions, text="Edit", command=lambda g=goal: edit_goal(g),
            bg=theme["card_alt"], fg=theme["text"],
            activebackground=theme["border"], relief="flat", bd=0,
            font=(SANS, 8, "bold"), padx=12, pady=6, cursor="hand2"
        )
        edit_btn.pack(side="left", padx=(0, 5))

        delete_btn = tk.Button(
            actions, text="Delete", command=lambda g=goal: delete_goal(g),
            bg=theme["card_alt"], fg=theme["danger"],
            activebackground=theme["border"], relief="flat", bd=0,
            font=(SANS, 8, "bold"), padx=12, pady=6, cursor="hand2"
        )
        delete_btn.pack(side="left")


# ---------- Analytics ----------

def month_label(y, m):
    return datetime(y, m, 1).strftime("%b")


def previous_month(y, m):
    m -= 1
    if m == 0:
        y -= 1
        m = 12
    return y, m


def show_analytics():
    show_page("Analytics", build_analytics)


def build_analytics():
    now = datetime.now()
    page_header(
        "Financial Intelligence",
        "A factual view of where your money has been moving.",
    )

    outer, inner = scrollable_frame(content)
    outer.pack(fill="both", expand=True, padx=35, pady=0)

    six_months = []
    y, m = now.year, now.month
    for _ in range(6):
        six_months.append((y, m))
        y, m = previous_month(y, m)
    six_months.reverse()

    chart_card = make_card(inner)
    chart_card.pack(fill="x", pady=5)
    section_title(chart_card, "Six-Month Movement", "Income vs spending")

    canvas = tk.Canvas(
        chart_card, height=240, bg=theme["card"],
        highlightthickness=0
    )
    canvas.pack(fill="x", padx=20, pady=(0, 16))

    values = []
    for yy, mm in six_months:
        inc, exp, net = monthly_totals(yy, mm)
        values.extend([inc, exp])

    max_value = max(values, default=1)
    max_value = max(max_value, 1)

    width = 830
    height = 200
    left_margin = 45
    bottom = 170
    usable_w = max(width - left_margin - 20, 100)
    group_w = usable_w / len(six_months)

    canvas.create_line(left_margin, 15, left_margin, bottom, fill=theme["border"])
    canvas.create_line(left_margin, bottom, width - 10, bottom, fill=theme["border"])

    for i, (yy, mm) in enumerate(six_months):
        inc, exp, net = monthly_totals(yy, mm)
        x_center = left_margin + group_w * i + group_w / 2
        bar_w = min(group_w * 0.23, 32)

        inc_h = (inc / max_value) * 135
        exp_h = (exp / max_value) * 135

        canvas.create_rectangle(
            x_center - bar_w - 2, bottom - inc_h,
            x_center - 2, bottom,
            fill=theme["success"], outline=""
        )
        canvas.create_rectangle(
            x_center + 2, bottom - exp_h,
            x_center + bar_w + 2, bottom,
            fill=theme["danger"], outline=""
        )
        canvas.create_text(
            x_center, bottom + 15, text=month_label(yy, mm),
            fill=theme["muted"], font=(SANS, 8)
        )

    canvas.create_text(
        20, 22, text=money(max_value),
        fill=theme["muted"], font=(SANS, 7), anchor="w"
    )
    canvas.create_text(
        width - 125, 20, text="IN", fill=theme["success"], font=(SANS, 8, "bold")
    )
    canvas.create_text(
        width - 90, 20, text="OUT", fill=theme["danger"], font=(SANS, 8, "bold")
    )

    breakdown = make_card(inner)
    breakdown.pack(fill="x", pady=5)
    section_title(breakdown, "Current Month", now.strftime("%B %Y"))

    inc, exp, net = monthly_totals(now.year, now.month)
    for title, value, color in [
        ("Income", inc, theme["success"]),
        ("Spending", exp, theme["danger"]),
        ("Net", net, theme["text"]),
    ]:
        row = tk.Frame(breakdown, bg=theme["card"])
        row.pack(fill="x", padx=20, pady=3)
        tk.Label(
            row, text=title, bg=theme["card"], fg=theme["muted"],
            font=(SANS, 8, "bold"), width=15, anchor="w"
        ).pack(side="left")
        tk.Label(
            row, text=money(value), bg=theme["card"], fg=color,
            font=(SANS, 10, "bold"), anchor="w"
        ).pack(side="left")

    add_divider(breakdown, 8)

    cat_totals = month_category_totals(now.year, now.month)
    if cat_totals:
        max_cat = max(cat_totals.values())
        for cat, value in sorted(cat_totals.items(), key=lambda x: x[1], reverse=True):
            row = tk.Frame(breakdown, bg=theme["card"])
            row.pack(fill="x", padx=20, pady=4)
            tk.Label(
                row, text=cat, bg=theme["card"], fg=theme["text"],
                font=(SANS, 8), width=17, anchor="w"
            ).pack(side="left")
            bar = tk.Frame(row, bg=theme["card_alt"], height=8)
            bar.pack(side="left", fill="x", expand=True, padx=5)
            bar.pack_propagate(False)
            fill = tk.Frame(bar, bg=theme["accent"], height=8)
            fill.place(relwidth=(value / max_cat if max_cat else 0), relheight=1)
            tk.Label(
                row, text=money(value), bg=theme["card"], fg=theme["muted"],
                font=(SANS, 8), width=14, anchor="e"
            ).pack(side="right")
    else:
        tk.Label(
            breakdown, text="No spending has been recorded this month.",
            bg=theme["card"], fg=theme["muted"], font=(SERIF, 10, "italic")
        ).pack(anchor="w", padx=20, pady=12)

    forecast = make_card(inner)
    forecast.pack(fill="x", pady=5)
    section_title(forecast, "Savings Pace", "Simple estimates, not guarantees")

    avg = average_monthly_savings()
    target_left = sum(max(g["target"] - g["saved"], 0) for g in goals if not g["completed"])

    tk.Label(
        forecast,
        text=f"Average monthly net over the last 3 months  •  {money(avg)}",
        bg=theme["card"], fg=theme["text"], font=(SANS, 9, "bold")
    ).pack(anchor="w", padx=20, pady=5)

    tk.Label(
        forecast,
        text=f"Remaining across active goals  •  {money(target_left)}",
        bg=theme["card"], fg=theme["muted"], font=(SANS, 9)
    ).pack(anchor="w", padx=20, pady=3)

    tk.Label(
        forecast,
        text="These figures are calculated from the entries currently stored in Vault.",
        bg=theme["card"], fg=theme["muted"], font=(SANS, 8)
    ).pack(anchor="w", padx=20, pady=(3, 14))


# ---------- Recurring ----------

def show_recurring():
    show_page("Recurring", build_recurring)


def recurring_dialog(existing=None):
    editing = existing is not None

    dialog = tk.Toplevel(window)
    dialog.title("Edit Recurring Entry" if editing else "New Recurring Entry")
    dialog.geometry("450x610")
    dialog.resizable(False, False)
    dialog.configure(bg=theme["bg"])
    dialog.transient(window)
    dialog.grab_set()

    tk.Label(
        dialog, text="Recurring Entry",
        bg=theme["bg"], fg=theme["text"],
        font=(SERIF_BOLD, 21, "bold")
    ).pack(pady=(20, 4))

    tk.Label(
        dialog, text="Vault will record it automatically when its date arrives.",
        bg=theme["bg"], fg=theme["muted"], font=(SANS, 9)
    ).pack(pady=(0, 12))

    form = tk.Frame(
        dialog, bg=theme["card"],
        highlightbackground=theme["border"], highlightthickness=1
    )
    form.pack(fill="x", padx=28)

    def entry(label, value=""):
        tk.Label(
            form, text=label, bg=theme["card"], fg=theme["muted"],
            font=(SANS, 8, "bold")
        ).pack(anchor="w", padx=18, pady=(9, 3))
        box = tk.Entry(
            form, bg=theme["input"], fg=theme["text"],
            insertbackground=theme["text"], relief="flat",
            highlightbackground=theme["border"], highlightthickness=1,
            font=(SANS, 9)
        )
        box.pack(fill="x", padx=18, ipady=5)
        if value:
            box.insert(0, value)
        return box

    name_entry = entry("NAME", existing["name"] if editing else "")
    amount_entry = entry("AMOUNT", str(existing["amount"]) if editing else "")
    note_entry = entry("NOTE", existing["note"] if editing else "Recurring entry")

    tk.Label(
        form, text="TYPE", bg=theme["card"], fg=theme["muted"],
        font=(SANS, 8, "bold")
    ).pack(anchor="w", padx=18, pady=(7, 3))
    type_var = tk.StringVar(
        value=("Income" if existing and existing["type"] == "income" else "Expense")
    )
    type_menu = ttk.Combobox(
        form, textvariable=type_var,
        values=["Income", "Expense"], state="readonly", font=(SANS, 9)
    )
    type_menu.pack(fill="x", padx=18, pady=(0, 5), ipady=4)

    tk.Label(
        form, text="FREQUENCY", bg=theme["card"], fg=theme["muted"],
        font=(SANS, 8, "bold")
    ).pack(anchor="w", padx=18, pady=(7, 3))
    frequency_var = tk.StringVar(
        value=existing["frequency"] if editing else "Monthly"
    )
    frequency_menu = ttk.Combobox(
        form, textvariable=frequency_var,
        values=["Daily", "Weekly", "Monthly", "Yearly"],
        state="readonly", font=(SANS, 9)
    )
    frequency_menu.pack(fill="x", padx=18, pady=(0, 5), ipady=4)

    tk.Label(
        form, text="NEXT DATE  (YYYY-MM-DD)", bg=theme["card"], fg=theme["muted"],
        font=(SANS, 8, "bold")
    ).pack(anchor="w", padx=18, pady=(7, 3))
    next_date_entry = tk.Entry(
        form, bg=theme["input"], fg=theme["text"],
        insertbackground=theme["text"], relief="flat",
        highlightbackground=theme["border"], highlightthickness=1,
        font=(SANS, 9)
    )
    next_date_entry.pack(fill="x", padx=18, ipady=5, pady=(0, 5))
    next_date_entry.insert(
        0, existing["next_date"] if editing else date.today().strftime("%Y-%m-%d")
    )

    tk.Label(
        form, text="CATEGORY", bg=theme["card"], fg=theme["muted"],
        font=(SANS, 8, "bold")
    ).pack(anchor="w", padx=18, pady=(7, 3))
    category_var = tk.StringVar(
        value=existing.get("category", "Other") if editing else "Other"
    )
    category_menu = ttk.Combobox(
        form, textvariable=category_var,
        values=categories or BUILT_IN_CATEGORIES,
        state="readonly", font=(SANS, 9)
    )
    category_menu.pack(fill="x", padx=18, pady=(0, 5), ipady=4)

    tk.Label(
        form, text="GOAL", bg=theme["card"], fg=theme["muted"],
        font=(SANS, 8, "bold")
    ).pack(anchor="w", padx=18, pady=(7, 3))
    goal_var = tk.StringVar(
        value=(existing.get("goal", "") or "None") if editing else "None"
    )
    goal_menu = ttk.Combobox(
        form, textvariable=goal_var,
        values=["None"] + goal_names(), state="readonly", font=(SANS, 9)
    )
    goal_menu.pack(fill="x", padx=18, pady=(0, 12), ipady=4)

    def submit():
        try:
            name = name_entry.get().strip()
            amount = parse_amount(amount_entry.get())
            next_date = next_date_entry.get().strip()
            validate_date(next_date)
            if not name or amount <= 0:
                raise ValueError

            item = {
                "name": name,
                "type": "income" if type_var.get() == "Income" else "expense",
                "amount": amount,
                "note": note_entry.get().strip() or "Recurring entry",
                "category": category_var.get() or "Other",
                "goal": "" if goal_var.get() == "None" else goal_var.get(),
                "frequency": frequency_var.get(),
                "next_date": next_date,
                "active": True if not editing else existing["active"],
            }

            if editing:
                item["id"] = existing["id"]
                update_recurring_db(item)
            else:
                save_recurring(item)

            dialog.destroy()
            reload_data_and_page()

        except ValueError:
            messagebox.showerror(
                "Invalid Recurring Entry",
                "Enter a name, positive amount, and valid next date.",
                parent=dialog
            )

    button = tk.Button(
        dialog,
        text="Save Recurring Entry",
        command=submit,
        bg=theme["accent"], fg="#fffaf0",
        activebackground=theme["accent_dark"], relief="flat", bd=0,
        font=(SANS, 10, "bold"), padx=24, pady=9, cursor="hand2"
    )
    button.pack(pady=12)
    hover_button(button, theme["accent"], theme["accent_dark"])


def add_recurring():
    recurring_dialog()


def edit_recurring(item):
    recurring_dialog(existing=item)


def toggle_recurring(item):
    item["active"] = not item["active"]
    update_recurring_db(item)
    reload_data_and_page()


def delete_recurring(item):
    if not messagebox.askyesno(
        "Delete Recurring Entry",
        f"Delete “{item['name']}”?",
        parent=window
    ):
        return
    delete_recurring_db(item["id"])
    reload_data_and_page()


def process_recurring():
    today = date.today()

    for item in list(recurring_items):
        if not item["active"]:
            continue

        try:
            next_date = datetime.strptime(item["next_date"], "%Y-%m-%d").date()
        except ValueError:
            continue

        changed = False
        while next_date <= today:
            entry = {
                "type": item["type"],
                "amount": item["amount"],
                "note": f"{item['note']}  (Recurring)",
                "date": next_date.strftime("%Y-%m-%d"),
                "time": datetime.now().strftime("%I:%M:%S %p"),
                "goal": item.get("goal", ""),
                "category": item.get("category", "Other"),
                "created_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            }

            current_balance = sum(transaction_effect(t) for t in transactions)
            if item["type"] == "expense" and item["amount"] > current_balance:
                # Avoid creating a debt entry. Advance the schedule so a failed
                # past occurrence does not get attempted forever on every launch.
                next_date = date_plus_frequency(
                    next_date.strftime("%Y-%m-%d"), item["frequency"]
                )
                changed = True
                continue

            new_id = save_transaction(entry)
            entry["id"] = new_id
            transactions.append(entry)

            if item.get("goal"):
                adjust_goal(item["goal"], transaction_effect(entry))

            next_date = date_plus_frequency(
                next_date.strftime("%Y-%m-%d"), item["frequency"]
            )
            changed = True

        if changed:
            item["next_date"] = next_date.strftime("%Y-%m-%d")
            update_recurring_db(item)


def build_recurring():
    header = page_header(
        "Recurring Entries",
        "Allow regular allowance, bills, or subscriptions to appear automatically.",
    )

    add_btn = tk.Button(
        header, text="＋ New Recurring", command=add_recurring,
        bg=theme["accent"], fg="#fffaf0",
        activebackground=theme["accent_dark"], relief="flat", bd=0,
        font=(SANS, 9, "bold"), padx=15, pady=8, cursor="hand2"
    )
    add_btn.pack(side="right", anchor="n")
    hover_button(add_btn, theme["accent"], theme["accent_dark"])

    outer, inner = scrollable_frame(content)
    outer.pack(fill="both", expand=True, padx=35, pady=3)

    if not recurring_items:
        card = make_card(inner)
        card.pack(fill="x")
        tk.Label(
            card, text="No recurring entries.",
            bg=theme["card"], fg=theme["muted"],
            font=(SERIF, 11, "italic")
        ).pack(pady=25)
        return

    for item in recurring_items:
        card = tk.Frame(
            inner, bg=theme["card"],
            highlightbackground=theme["border"], highlightthickness=1
        )
        card.pack(fill="x", pady=5)

        title_row = tk.Frame(card, bg=theme["card"])
        title_row.pack(fill="x", padx=20, pady=(13, 4))

        color = theme["success"] if item["type"] == "income" else theme["danger"]

        tk.Label(
            title_row, text=item["name"],
            bg=theme["card"], fg=theme["text"],
            font=(SERIF_BOLD, 13, "bold")
        ).pack(side="left")

        tk.Label(
            title_row, text=f"{'+' if item['type']=='income' else '−'} {money(item['amount'])}",
            bg=theme["card"], fg=color,
            font=(SANS, 10, "bold")
        ).pack(side="right")

        tk.Label(
            card,
            text=f"{item['frequency']}  •  Next: {item['next_date']}  •  "
                 f"{item.get('category','Other')}"
                 + (f"  •  Goal: {item['goal']}" if item.get("goal") else ""),
            bg=theme["card"], fg=theme["muted"], font=(SANS, 8)
        ).pack(anchor="w", padx=20)

        tk.Label(
            card, text=item["note"],
            bg=theme["card"], fg=theme["text"], font=(SANS, 8)
        ).pack(anchor="w", padx=20, pady=(4, 0))

        actions = tk.Frame(card, bg=theme["card"])
        actions.pack(fill="x", padx=20, pady=(8, 12))

        toggle = tk.Button(
            actions, text="Pause" if item["active"] else "Resume",
            command=lambda x=item: toggle_recurring(x),
            bg=theme["card_alt"], fg=theme["text"],
            activebackground=theme["border"], relief="flat", bd=0,
            font=(SANS, 8, "bold"), padx=11, pady=6, cursor="hand2"
        )
        toggle.pack(side="left", padx=(0, 5))

        edit = tk.Button(
            actions, text="Edit", command=lambda x=item: edit_recurring(x),
            bg=theme["card_alt"], fg=theme["text"],
            activebackground=theme["border"], relief="flat", bd=0,
            font=(SANS, 8, "bold"), padx=11, pady=6, cursor="hand2"
        )
        edit.pack(side="left", padx=(0, 5))

        delete = tk.Button(
            actions, text="Delete", command=lambda x=item: delete_recurring(x),
            bg=theme["card_alt"], fg=theme["danger"],
            activebackground=theme["border"], relief="flat", bd=0,
            font=(SANS, 8, "bold"), padx=11, pady=6, cursor="hand2"
        )
        delete.pack(side="left")


# ---------- Settings / Data ----------

def export_csv():
    path = filedialog.asksaveasfilename(
        parent=window,
        defaultextension=".csv",
        filetypes=[("CSV file", "*.csv")],
        initialfile=f"vault_export_{datetime.now().strftime('%Y%m%d')}.csv"
    )
    if not path:
        return

    try:
        with open(path, "w", newline="", encoding="utf-8-sig") as handle:
            writer = csv.writer(handle)
            writer.writerow([
                "id", "type", "amount", "note", "date", "time",
                "goal", "category"
            ])
            for t in transactions:
                writer.writerow([
                    t["id"], t["type"], t["amount"], t["note"], t["date"],
                    t["time"], t.get("goal", ""), t.get("category", "Other")
                ])
        messagebox.showinfo("Export Complete", f"CSV saved to:\n{path}", parent=window)
    except Exception as exc:
        messagebox.showerror("Export Failed", str(exc), parent=window)


def export_xlsx():
    path = filedialog.asksaveasfilename(
        parent=window,
        defaultextension=".xlsx",
        filetypes=[("Excel workbook", "*.xlsx")],
        initialfile=f"vault_export_{datetime.now().strftime('%Y%m%d')}.xlsx"
    )
    if not path:
        return

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Transactions"

        headers = [
            "ID", "Type", "Amount", "Note", "Date", "Time", "Goal", "Category"
        ]
        sheet.append(headers)

        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")

        for t in transactions:
            sheet.append([
                t["id"], t["type"], t["amount"], t["note"], t["date"],
                t["time"], t.get("goal", ""), t.get("category", "Other")
            ])

        for column, width in {
            "A": 8, "B": 12, "C": 15, "D": 35,
            "E": 15, "F": 15, "G": 25, "H": 18
        }.items():
            sheet.column_dimensions[column].width = width

        goals_sheet = workbook.create_sheet("Goals")
        goals_sheet.append([
            "ID", "Name", "Target", "Saved", "Remaining",
            "Deadline", "Description", "Completed"
        ])
        for g in goals:
            goals_sheet.append([
                g["id"], g["name"], g["target"], g["saved"],
                max(g["target"] - g["saved"], 0), g.get("deadline", ""),
                g.get("description", ""), "Yes" if g["completed"] else "No"
            ])

        for cell in goals_sheet[1]:
            cell.font = Font(bold=True)

        workbook.save(path)
        messagebox.showinfo("Export Complete", f"Excel file saved to:\n{path}", parent=window)
    except ImportError:
        messagebox.showerror(
            "Excel Export Unavailable",
            "The optional openpyxl package is not installed.\n\n"
            "CSV export is still available.",
            parent=window
        )
    except Exception as exc:
        messagebox.showerror("Export Failed", str(exc), parent=window)


def import_csv():
    path = filedialog.askopenfilename(
        parent=window,
        filetypes=[("CSV file", "*.csv")]
    )
    if not path:
        return

    try:
        with open(path, "r", newline="", encoding="utf-8-sig") as handle:
            reader = csv.DictReader(handle)
            required = {"type", "amount", "note", "date"}
            if not required.issubset(set(reader.fieldnames or [])):
                raise ValueError(
                    "The CSV must contain at least type, amount, note, and date columns."
                )

            imported = 0
            with get_db() as db:
                for row in reader:
                    row_type = row["type"].strip().lower()
                    if row_type not in ("income", "expense"):
                        continue
                    amount = parse_amount(row["amount"])
                    if amount <= 0:
                        continue
                    date_text = row["date"].strip()
                    validate_date(date_text)
                    time_text = row.get("time", "").strip() or "12:00:00 AM"
                    category = row.get("category", "").strip() or "Other"
                    goal = row.get("goal", "").strip()
                    note = row.get("note", "").strip() or "Imported entry"

                    db.execute("""
                        INSERT INTO transactions
                            (type, amount, note, date, time, goal, category, created_at)
                        VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                    """, (
                        row_type, amount, note, date_text, time_text,
                        goal, category, datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                    ))
                    imported += 1

        reload_data_and_page()
        messagebox.showinfo(
            "Import Complete",
            f"Imported {imported} transaction(s).\n\n"
            "Imported entries do not alter goal allocations automatically.",
            parent=window
        )
    except Exception as exc:
        messagebox.showerror("Import Failed", str(exc), parent=window)


def create_manual_backup():
    try:
        BACKUP_DIR.mkdir(parents=True, exist_ok=True)
        with get_db() as db:
            db.execute("PRAGMA wal_checkpoint(FULL)")
        target = BACKUP_DIR / f"vault_{datetime.now().strftime('%Y-%m-%d_%H%M%S')}.db"
        shutil.copy2(DB_PATH, target)
        messagebox.showinfo("Backup Created", f"Backup saved to:\n{target}", parent=window)
    except Exception as exc:
        messagebox.showerror("Backup Failed", str(exc), parent=window)


def run_integrity_check():
    try:
        result = integrity_check()
        if result == "ok":
            messagebox.showinfo(
                "Database Integrity",
                "Vault's database integrity check passed.",
                parent=window
            )
        else:
            messagebox.showerror(
                "Database Integrity",
                f"SQLite reported:\n{result}",
                parent=window
            )
    except Exception as exc:
        messagebox.showerror("Integrity Check Failed", str(exc), parent=window)


def add_custom_category():
    dialog = tk.Toplevel(window)
    dialog.title("New Category")
    dialog.geometry("360x220")
    dialog.resizable(False, False)
    dialog.configure(bg=theme["bg"])
    dialog.transient(window)
    dialog.grab_set()

    tk.Label(
        dialog, text="New Category",
        bg=theme["bg"], fg=theme["text"],
        font=(SERIF_BOLD, 19, "bold")
    ).pack(pady=(20, 4))

    entry = tk.Entry(
        dialog, bg=theme["input"], fg=theme["text"],
        insertbackground=theme["text"], relief="flat",
        highlightbackground=theme["border"], highlightthickness=1,
        font=(SANS, 10)
    )
    entry.pack(fill="x", padx=28, ipady=7, pady=12)

    def submit():
        name = entry.get().strip()
        if not name:
            return
        try:
            add_category_db(name)
            dialog.destroy()
            reload_data_and_page()
        except sqlite3.IntegrityError:
            messagebox.showerror("Duplicate Category", "That category already exists.", parent=dialog)

    btn = tk.Button(
        dialog, text="Create Category", command=submit,
        bg=theme["accent"], fg="#fffaf0",
        activebackground=theme["accent_dark"],
        relief="flat", bd=0, font=(SANS, 9, "bold"),
        padx=18, pady=8, cursor="hand2"
    )
    btn.pack()
    hover_button(btn, theme["accent"], theme["accent_dark"])
    entry.focus_set()
    dialog.bind("<Return>", lambda e: submit())
    dialog.bind("<Escape>", lambda e: dialog.destroy())


def remove_custom_category():
    custom = [c for c in categories if c not in BUILT_IN_CATEGORIES]
    if not custom:
        messagebox.showinfo(
            "Categories",
            "There are no custom categories to remove.",
            parent=window
        )
        return

    dialog = tk.Toplevel(window)
    dialog.title("Remove Category")
    dialog.geometry("360x240")
    dialog.resizable(False, False)
    dialog.configure(bg=theme["bg"])
    dialog.transient(window)
    dialog.grab_set()

    tk.Label(
        dialog, text="Remove Custom Category",
        bg=theme["bg"], fg=theme["text"],
        font=(SERIF_BOLD, 18, "bold")
    ).pack(pady=(20, 8))

    var = tk.StringVar(value=custom[0])
    menu = ttk.Combobox(
        dialog, textvariable=var, values=custom,
        state="readonly", width=24
    )
    menu.pack(pady=8)

    def submit():
        delete_category_db(var.get())
        dialog.destroy()
        reload_data_and_page()

    btn = tk.Button(
        dialog, text="Remove", command=submit,
        bg=theme["danger"], fg="#fffaf0",
        activebackground="#74342d", relief="flat", bd=0,
        font=(SANS, 9, "bold"), padx=18, pady=8, cursor="hand2"
    )
    btn.pack(pady=8)


# ---------- Settings ----------

def toggle_dark_mode():
    global dark_mode, theme
    dark_mode = not dark_mode
    theme = DARK if dark_mode else LIGHT
    save_dark_mode()
    rebuild_shell()


def show_settings():
    show_page("Settings", build_settings)


def build_settings():
    page_header(
        "Settings & Archive",
        "Shape Vault to your own taste and protect its records.",
    )

    outer, inner = scrollable_frame(content)
    outer.pack(fill="both", expand=True, padx=35, pady=0)

    appearance = make_card(inner)
    appearance.pack(fill="x", pady=5)
    section_title(appearance, "Appearance")

    mode_row = tk.Frame(appearance, bg=theme["card"])
    mode_row.pack(fill="x", padx=20, pady=(0, 17))

    tk.Label(
        mode_row, text="Dark mode",
        bg=theme["card"], fg=theme["text"],
        font=(SERIF_BOLD, 13, "bold")
    ).pack(side="left")

    tk.Label(
        mode_row, text="The entire interface follows this theme.",
        bg=theme["card"], fg=theme["muted"], font=(SANS, 9)
    ).pack(side="left", padx=15)

    button = tk.Button(
        mode_row, text="ON" if dark_mode else "OFF",
        command=toggle_dark_mode,
        bg=theme["accent"] if dark_mode else theme["card_alt"],
        fg="#fffaf0" if dark_mode else theme["text"],
        activebackground=theme["accent_dark"],
        relief="flat", bd=0, font=(SANS, 9, "bold"),
        width=8, cursor="hand2"
    )
    button.pack(side="right")

    categories_card = make_card(inner)
    categories_card.pack(fill="x", pady=5)
    section_title(categories_card, "Categories", "Use them to understand spending")

    tk.Label(
        categories_card,
        text=", ".join(categories),
        bg=theme["card"], fg=theme["muted"],
        font=(SANS, 8), wraplength=820, justify="left"
    ).pack(anchor="w", padx=20, pady=(0, 10))

    category_actions = tk.Frame(categories_card, bg=theme["card"])
    category_actions.pack(anchor="w", padx=20, pady=(0, 15))

    add_cat = tk.Button(
        category_actions, text="＋ Add Category",
        command=add_custom_category,
        bg=theme["card_alt"], fg=theme["text"],
        activebackground=theme["border"], relief="flat", bd=0,
        font=(SANS, 8, "bold"), padx=12, pady=6, cursor="hand2"
    )
    add_cat.pack(side="left", padx=(0, 5))

    rem_cat = tk.Button(
        category_actions, text="Remove Custom",
        command=remove_custom_category,
        bg=theme["card_alt"], fg=theme["danger"],
        activebackground=theme["border"], relief="flat", bd=0,
        font=(SANS, 8, "bold"), padx=12, pady=6, cursor="hand2"
    )
    rem_cat.pack(side="left")

    archive = make_card(inner)
    archive.pack(fill="x", pady=5)
    section_title(archive, "Data & Archive")

    info = tk.Frame(archive, bg=theme["card"])
    info.pack(fill="x", padx=20, pady=(0, 10))

    rows = [
        ("Database", str(DB_PATH)),
        ("Automatic backup", last_backup_text()),
        ("Entries", str(len(transactions))),
        ("Goals", str(len(goals))),
    ]
    for label, value in rows:
        row = tk.Frame(info, bg=theme["card"])
        row.pack(fill="x", pady=2)
        tk.Label(
            row, text=label, width=18, anchor="w",
            bg=theme["card"], fg=theme["muted"], font=(SANS, 8, "bold")
        ).pack(side="left")
        tk.Label(
            row, text=value, bg=theme["card"], fg=theme["text"],
            font=(SANS, 8), anchor="w"
        ).pack(side="left", fill="x", expand=True)

    actions = tk.Frame(archive, bg=theme["card"])
    actions.pack(fill="x", padx=20, pady=(4, 16))

    buttons = [
        ("Backup Now", create_manual_backup, theme["accent"], theme["accent_dark"]),
        ("Export CSV", export_csv, theme["card_alt"], theme["border"]),
        ("Export Excel", export_xlsx, theme["card_alt"], theme["border"]),
        ("Import CSV", import_csv, theme["card_alt"], theme["border"]),
        ("Check Database", run_integrity_check, theme["card_alt"], theme["border"]),
    ]

    for i, (label, command, bg, hover) in enumerate(buttons):
        b = tk.Button(
            actions, text=label, command=command,
            bg=bg, fg="#fffaf0" if bg == theme["accent"] else theme["text"],
            activebackground=hover, relief="flat", bd=0,
            font=(SANS, 8, "bold"), padx=10, pady=7, cursor="hand2"
        )
        b.pack(side="left", padx=(0, 5))
        hover_button(b, bg, hover)

    notes = make_card(inner)
    notes.pack(fill="x", pady=5)
    section_title(notes, "Useful Shortcuts")

    shortcut_text = (
        "Ctrl + N   Add money\n"
        "Ctrl + F   Focus transaction search\n"
        "Ctrl + G   Open goals\n"
        "Ctrl + Z   Undo the last transaction change\n"
        "Escape     Close a dialog"
    )
    tk.Label(
        notes, text=shortcut_text, bg=theme["card"], fg=theme["muted"],
        font=(SANS, 9), justify="left"
    ).pack(anchor="w", padx=20, pady=(0, 18))


# ---------- Navigation ----------

def nav_button(text, command):
    button = tk.Button(
        sidebar, text=text, command=command,
        bg=theme["sidebar"], fg=theme["sidebar_text"],
        activebackground=theme["sidebar_hover"],
        activeforeground=theme["sidebar_text"],
        relief="flat", anchor="w", padx=25,
        font=(SANS, 9), cursor="hand2", bd=0
    )
    button.pack(fill="x", ipady=10)
    hover_button(button, theme["sidebar"], theme["sidebar_hover"])
    return button


def rebuild_shell():
    global sidebar, content, clock_label

    for widget in window.winfo_children():
        widget.destroy()

    sidebar = tk.Frame(window, bg=theme["sidebar"], width=220)
    sidebar.pack(side="left", fill="y")
    sidebar.pack_propagate(False)

    tk.Label(
        sidebar, text="V A U L T",
        bg=theme["sidebar"], fg=theme["accent"],
        font=(SERIF_BOLD, 22, "bold")
    ).pack(pady=(29, 3))

    tk.Label(
        sidebar, text="PERSONAL LEDGER",
        bg=theme["sidebar"], fg=theme["sidebar_muted"],
        font=(SANS, 7, "bold")
    ).pack(pady=(0, 24))

    add_divider(sidebar, 0)

    nav_button("  ⌂   Dashboard", show_dashboard)
    nav_button("  ≡   Transactions", show_transactions)
    nav_button("  □   Calendar", show_calendar)
    nav_button("  ◈   Goals", show_goals)
    nav_button("  ◉   Analytics", show_analytics)
    nav_button("  ↻   Recurring", show_recurring)
    nav_button("  ⚙   Settings", show_settings)

    tk.Frame(sidebar, bg=theme["sidebar"]).pack(fill="both", expand=True)

    tk.Label(
        sidebar, text="✦  MEMENTO VIVERE  ✦",
        bg=theme["sidebar"], fg=theme["sidebar_muted"],
        font=(SERIF, 8, "italic")
    ).pack(pady=(0, 5))

    tk.Label(
        sidebar, text="Keep account of what matters.",
        bg=theme["sidebar"], fg=theme["sidebar_muted"],
        font=(SANS, 7)
    ).pack(pady=(0, 20))

    main = tk.Frame(window, bg=theme["bg"])
    main.pack(side="right", expand=True, fill="both")

    clock_strip = tk.Frame(main, bg=theme["bg"])
    clock_strip.pack(fill="x", padx=35, pady=(12, 0))

    clock_label = tk.Label(
        clock_strip, text="", bg=theme["bg"], fg=theme["muted"],
        font=(SANS, 8)
    )
    clock_label.pack(side="right")

    content = tk.Frame(main, bg=theme["bg"])
    content.pack(side="bottom", expand=True, fill="both")

    refresh_current_page()


# ---------- Shortcuts ----------

def focus_transaction_search():
    if current_page == "Transactions":
        for widget in content.winfo_children():
            # Search field is the first entry in the toolbar.
            entries = []
            def collect(w):
                if isinstance(w, tk.Entry):
                    entries.append(w)
                for child in w.winfo_children():
                    collect(child)
            collect(content)
            if entries:
                entries[0].focus_set()
                entries[0].selection_range(0, tk.END)
                break
    else:
        show_transactions()


window.bind("<Control-n>", lambda e: add_money())
window.bind("<Control-f>", lambda e: focus_transaction_search())
window.bind("<Control-g>", lambda e: show_goals())
window.bind("<Control-z>", lambda e: undo_last())
window.bind("<Escape>", lambda e: None)


# ---------- Start ----------

init_db()
load_data()
process_recurring()
load_data()
backup_database()

window.configure(bg=theme["bg"])
rebuild_shell()
update_clock()
window.mainloop()
